import os
import requests
import json,logging,time
from io import BytesIO
from pypdf import PdfReader
from docx import Document
from datetime import datetime,timedelta,timezone
from pathlib import Path
from django.conf import settings
from django.db.models import F
from django.http import JsonResponse, HttpResponse
from django.shortcuts import render
from django.views.decorators.csrf import ensure_csrf_cookie
from . import claude,demo,matcher
from .models import UserWorkspace, DailyAIUsage
from django.utils import timezone as django_timezone
log=logging.getLogger(__name__)
@ensure_csrf_cookie
def index(request):
    limit = max(0, int(os.getenv('HIREATIONSHIP_DAILY_GENERATION_LIMIT', '4')))
    usage = DailyAIUsage.objects.filter(user=request.user, day=django_timezone.localdate()).first()
    used = usage.generations if usage else 0
    from .job_details import advert_read_capability
    response = render(request, 'index.html', {
        **advert_read_capability(),
        'generation_limit': limit,
        'generation_remaining': max(0, limit - used),
    })
    response['Cache-Control'] = 'no-store'
    return response
def endpoint(method='POST'):
 def decorate(fn):
  def wrapped(request):
   if request.method!=method:return JsonResponse({'error':'Method not allowed'},status=405)
   try:
    data=json.loads(request.body) if request.body else {}
    if not isinstance(data,dict):return JsonResponse({'error':'Expected a JSON object'},status=400)
    return JsonResponse(fn(data))
   except (ValueError,TypeError,KeyError) as exc:return JsonResponse({'error':str(exc)},status=400)
   except Exception:
    log.exception('API failure');return JsonResponse({'error':'Request failed. Check server logs.'},status=500)
  return wrapped
 return decorate
@endpoint('GET')
def status(data):
 from .job_details import advert_read_capability
 return {**advert_read_capability(),'mode':'claude' if claude.has_key() else 'demo', 'serverKeyAvailable':bool(os.getenv('ANTHROPIC_API_KEY')) and os.getenv('HIREATIONSHIP_ALLOW_SHARED_CLAUDE','0')=='1'}
@endpoint()
def profile(data):
 cv=data.get('cvText') or ''
 if len(cv.strip())<30:raise ValueError('Paste your CV text first.')
 return {'profile':claude.extract_profile(cv) if claude.has_key() else demo.demo_profile(cv),'mode':'claude' if claude.has_key() else 'demo'}
def load_roles():
 now=datetime.now(timezone.utc)
 raw=json.loads((settings.BASE_DIR/'data/roles.json').read_text())
 return [{**r,'sample':True,'closes':(now+timedelta(days=r['closesInDays'])).date().isoformat()} for r in raw]
@endpoint()
def match(data):
 p=data.get('profile') or {}
 if not p.get('gradYear'):raise ValueError('Set your graduation year first.')
 families=claude.suggest_families(p) if claude.has_key() else demo.demo_families(p)
 ids={f['id'] for f in families}
 roles=(data.get('extraRoles') or [])+load_roles()
 evaluated=matcher.match_all(p,roles)
 results=matcher.rank(evaluated)
 excluded=[{'title':r['role'].get('title','Role'),'reasons':r['hardFails']+(['Missing required skills: '+', '.join(r['missingRequired'])] if r['missingRequired'] else [])} for r in evaluated if r['tier']!='eligible']
 return {'stage':matcher.stage_for(p['gradYear']),'families':families,'results':[{**r,'inFamily':r['role'].get('sector') in ids or r['role'].get('custom') is True} for r in results],'hiddenCount':len(roles)-len(results),'excluded':excluded}
@endpoint()
def jd(data):
 if not claude.has_key():raise ValueError('Adding a role from a job description needs an Anthropic API key (see README).')
 text=data.get('text') or ''
 if len(text.strip())<80:raise ValueError('Paste the full job description.')
 role=claude.extract_role(text)
 return {'role':{**role,'id':'jd-'+str(int(time.time()*1000)),'custom':True,'sample':False,'stages':role.get('stages') or ['final','penultimate','early']}}
@endpoint()
def generate(data):
 p=data.get('profile');r=data.get('role')
 if not p or not r:raise ValueError('Missing profile or role.')
 ev=matcher.evaluate(p,r)
 return {**claude.generate_pack(p,r,ev),'mode':'claude'} if claude.has_key() else demo.demo_pack(p,r,ev)


def upload_cv(request):
    if request.method != 'POST':
        return JsonResponse({'error': 'Method not allowed'}, status=405)
    uploaded = request.FILES.get('cv')
    if not uploaded:
        return JsonResponse({'error': 'Choose a PDF or .docx file.'}, status=400)
    if uploaded.size > 5 * 1024 * 1024:
        return JsonResponse({'error': 'CV must be smaller than 5 MB.'}, status=400)
    extension = Path(uploaded.name).suffix.lower()
    if extension not in ('.pdf', '.docx'):
        return JsonResponse({'error': 'Only PDF and Word (.docx) files are supported.'}, status=400)
    try:
        content = uploaded.read()
        if extension == '.pdf':
            if not content.startswith(b'%PDF-'):
                raise ValueError('This does not appear to be a valid PDF.')
            reader = PdfReader(BytesIO(content))
            if len(reader.pages) > 25:
                raise ValueError('Please upload a CV of 25 pages or fewer.')
            cv_text = '\n'.join(page.extract_text() or '' for page in reader.pages)
        else:
            if not content.startswith(b'PK'):
                raise ValueError('This does not appear to be a valid Word document.')
            doc = Document(BytesIO(content))
            parts = [para.text for para in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(' | '.join(cell.text for cell in row.cells))
            cv_text = '\n'.join(parts)
        cv_text = cv_text.strip()
        if len(cv_text) < 30:
            raise ValueError('Could not extract enough text. For scanned PDFs, paste the CV text manually.')
        if len(cv_text) > 50000:
            raise ValueError('CV text is too long. Please shorten it to 50,000 characters.')
        return JsonResponse({'cvText': cv_text})
    except (ValueError, Exception) as exc:
        log.warning('CV upload failed: %s', type(exc).__name__)
        return JsonResponse({'error': str(exc) if isinstance(exc, ValueError) else 'Could not read this CV file.'}, status=400)


def export_pdf(request):
    """Render user-reviewed application text as a downloadable PDF."""
    if request.method != 'POST':
        return JsonResponse({'error': 'Method not allowed'}, status=405)
    try:
        from io import BytesIO
        from xml.sax.saxutils import escape
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        payload = json.loads(request.body)
        kind = payload.get('kind')
        content = payload.get('text', '')
        if kind not in ('cv', 'letter') or not isinstance(content, str) or not content.strip() or len(content) > 50000:
            raise ValueError('Choose a CV or cover letter containing up to 50,000 characters.')
        template = payload.get('template', 'classic')
        if template not in ('classic', 'modern'):
            raise ValueError('Invalid template.')
        styles = getSampleStyleSheet()
        navy = colors.HexColor('#263A67')
        accent = colors.HexColor('#E95B65') if template == 'modern' else navy
        body = ParagraphStyle('AppBody', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=14, spaceAfter=5, textColor=colors.HexColor('#222222'))
        heading = ParagraphStyle('AppHeading', parent=body, fontName='Helvetica-Bold', fontSize=11, leading=15, spaceBefore=10, spaceAfter=5, textColor=accent)
        title = ParagraphStyle('AppTitle', parent=heading, fontSize=16, leading=20, spaceAfter=12)
        bullet = ParagraphStyle('AppBullet', parent=body, leftIndent=14, firstLineIndent=-8)
        out = BytesIO()
        doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=48, rightMargin=48, topMargin=45, bottomMargin=45, title='Tailored CV' if kind == 'cv' else 'Cover letter')
        story = []
        for i, line in enumerate(content.splitlines()):
            line = line.strip()
            if not line:
                story.append(Spacer(1, 7))
                continue
            line = line.lstrip('#').strip()
            if line.startswith(('- ', '* ', '• ')):
                story.append(Paragraph('&#8226; ' + escape(line[2:]), bullet))
            elif kind == 'cv' and (i == 0 or (len(line) < 65 and (line.isupper() or line.endswith(':')))):
                story.append(Paragraph(escape(line.rstrip(':')), title if i == 0 else heading))
            else:
                story.append(Paragraph(escape(line), body))
        doc.build(story)
        response = HttpResponse(out.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="hireationship-' + ('cv' if kind == 'cv' else 'cover-letter') + '.pdf"'
        response['Cache-Control'] = 'no-store'
        return response
    except (ValueError, TypeError, json.JSONDecodeError) as exc:
        return JsonResponse({'error': str(exc)}, status=400)
    except Exception:
        log.exception('PDF export failed')
        return JsonResponse({'error': 'Could not generate PDF.'}, status=500)

def reserve_live_search(user, limit, amount=1):
    usage, _ = DailyAIUsage.objects.get_or_create(user=user, day=django_timezone.localdate())
    return DailyAIUsage.objects.filter(pk=usage.pk, job_searches__lte=limit-amount).update(job_searches=F('job_searches') + amount) == 1


def live_jobs(request):
    if request.method != 'POST':
        return JsonResponse({'error':'Method not allowed'},status=405)
    try:
        data = json.loads(request.body)
        profile = data.get('profile') if isinstance(data, dict) else None
        if not isinstance(profile, dict) or not profile.get('gradYear'):
            raise ValueError('Review your CV profile and graduation year before searching.')
        matcher.stage_for(profile['gradYear'])
    except (ValueError, TypeError, AttributeError):
        return JsonResponse({'error':'Review your CV profile and graduation year before searching.'},status=400)
    app_id=os.environ.get('ADZUNA_APP_ID')
    app_key=os.environ.get('ADZUNA_APP_KEY')
    if not app_id or not app_key:
        return JsonResponse({'error':'Live jobs are temporarily unavailable. The site administrator needs to configure the shared Adzuna connection.'},status=503)
    from .search_preferences import search_plan
    try:
        plans=search_plan(profile) if 'preferences' in profile else [{'what':str(data.get('what','graduate')).strip()[:90], 'where':str(data.get('where','')).strip()[:90]}]
    except (ValueError,TypeError,KeyError):
        return JsonResponse({'error':'Review your role, opportunity type and location preferences in 1 · You.'},status=400)
    query=plans[0]['what']
    where=str(data.get('where','London')).strip()[:90]
    if not query:
        return JsonResponse({'error':'Enter search keywords.'},status=400)
    try:
        page=int(data.get('page',1))
        if not 1 <= page <= 100: raise ValueError()
    except (TypeError,ValueError):
        return JsonResponse({'error':'Choose a search page between 1 and 100.'},status=400)
    limit = max(0, int(os.getenv('HIREATIONSHIP_DAILY_JOB_SEARCH_LIMIT', '30')))
    reserved = reserve_live_search(request.user, limit, len(plans))
    if not reserved:
        return JsonResponse({'error':f'Daily live-search limit reached ({limit}). Try again tomorrow. Demo matching is still available.'}, status=429)
    try:
        page_size = max(1, (15 if claude.has_key() else 50) // len(plans))
        raw_jobs=[]; counts=[]; has_more=False
        for plan in plans:
            response=requests.get(f'https://api.adzuna.com/v1/api/jobs/gb/search/{page}',params={'app_id':app_id,'app_key':app_key,'what':plan['what'],'where':plan['where'],'results_per_page':page_size,'sort_by':'date'},timeout=8)
            response.raise_for_status()
            payload=response.json()
            batch=payload.get('results',[])
            raw_jobs.extend(batch)
            count=payload.get('count')
            counts.append(count)
            if len(batch)==page_size and (not isinstance(count,int) or page*page_size<count): has_more=True
        # Different interests and locations can return the same listing.
        unique_jobs=list({str(item.get('id')):item for item in raw_jobs}.values())
        jobs=[]
        for item in unique_jobs:
            url=item.get('redirect_url','')
            if not url.startswith('https://'):
                continue
            jobs.append({'id':'adzuna-'+str(item.get('id','')),'title':str(item.get('title',''))[:200],
              'employer':str((item.get('company') or {}).get('display_name',''))[:150],
              'location':str((item.get('location') or {}).get('display_name',''))[:150],
              'summary':str(item.get('description',''))[:12000],
              'applyUrl':url,'posted':str(item.get('created',''))[:10],
              'salaryMin':item.get('salary_min'), 'salaryMax':item.get('salary_max')})
        assessments = (claude.assess_live_jobs(profile, jobs) if claude.has_key() else demo.assess_live_jobs(profile, jobs)) if jobs else []
        by_id = {a.get('id'): a for a in assessments if isinstance(a, dict)}
        matches = []
        review_jobs = []
        unmet_count = 0
        for job in jobs:
            assessment = by_id.get(job['id'], {})
            if assessment.get('eligible') is False:
                unmet_count += 1
            if assessment.get('eligible') is not True or not assessment.get('reasons') or not isinstance(assessment.get('role'), dict):
                if assessment.get('eligible') is not False:
                    review_jobs.append({'role': {**job, 'source':'Adzuna', 'sample':False, 'custom':True, 'eligibilityUnverified':True},
                        'tier':'unverified', 'category':'unverified', 'fit':None, 'hardPasses':[], 'matched':[], 'missingRequired':[], 'missingPreferred':[]})
                continue
            extracted = assessment['role']
            if not all(key in extracted for key in ('stages', 'minClassification', 'subjects', 'rightToWork', 'skillsRequired', 'skillsPreferred')):
                continue
            if not all(isinstance(extracted[key], list) for key in ('stages', 'subjects', 'skillsRequired', 'skillsPreferred')):
                continue
            if not isinstance(assessment['reasons'], list):
                continue
            role = {**extracted, **job, 'source': 'Adzuna', 'sample': False, 'custom': True, 'eligibilityChecked': True, 'matchingMode': 'claude' if claude.has_key() else 'rules'}
            evaluated = matcher.evaluate(profile, role)
            if evaluated['tier'] == 'eligible':
                evaluated['hardPasses'].extend(assessment['reasons'])
                matches.append(evaluated)
            else:
                unmet_count += 1
        from .job_details import advert_read_capability
        result=JsonResponse({**advert_read_capability(),'jobs':matcher.rank(matches, 'fit'), 'hiddenCount':len(jobs)-len(matches),
            'reviewJobs':review_jobs, 'fetchedCount':len(raw_jobs), 'totalAvailable':counts[0] if len(plans)==1 else None, 'queryCount':len(plans),
            'unmetCount':unmet_count, 'unclearCount':len(jobs)-len(matches)-unmet_count,
            'skippedCount':len(unique_jobs)-len(jobs), 'page':page, 'pageSize':page_size,
            'nextPage':page+1 if has_more and page<100 else None})
        result['Cache-Control']='no-store'
        return result
    except (requests.RequestException,ValueError):
        log.warning('Live job search failed; upstream request details omitted to protect API credentials.')
        return JsonResponse({'error':'Job search unavailable. Check credentials and PythonAnywhere outbound API access.'},status=502)

WORKSPACE_KEYS = ('cv', 'profile', 'parsed', 'apps', 'passed', 'passedHistory', 'extraRoles', 'jobMode')

def workspace(request):
    if request.method == 'GET':
        row = UserWorkspace.objects.filter(user=request.user).first()
        response = JsonResponse({'data': row.data if row else {}})
    elif request.method == 'POST':
        if len(request.body) > 500_000:
            return JsonResponse({'error': 'Workspace exceeds 500 KB. Remove old applications to free space.'}, status=413)
        try:
            payload = json.loads(request.body)
            if not isinstance(payload, dict) or not all(k in WORKSPACE_KEYS for k in payload):
                raise ValueError('Invalid workspace fields.')
            row, _ = UserWorkspace.objects.get_or_create(user=request.user)
            row.data = {**row.data, **payload}
            row.save(update_fields=['data', 'updated_at'])
            response = JsonResponse({'saved': True})
        except (ValueError, TypeError):
            return JsonResponse({'error': 'Invalid workspace data.'}, status=400)
    elif request.method == 'DELETE':
        UserWorkspace.objects.filter(user=request.user).delete()
        response = JsonResponse({'deleted': True})
    else:
        return JsonResponse({'error': 'Method not allowed'}, status=405)
    response['Cache-Control'] = 'no-store'
    return response

def generation_quota(request):
    if request.method != 'GET':
        return JsonResponse({'error': 'Method not allowed'}, status=405)
    limit = max(0, int(os.getenv('HIREATIONSHIP_DAILY_GENERATION_LIMIT', '4')))
    usage = DailyAIUsage.objects.filter(user=request.user, day=django_timezone.localdate()).first()
    used = usage.generations if usage else 0
    response = JsonResponse({'plan': 'free', 'limit': limit, 'used': used, 'remaining': max(0, limit-used)})
    response['Cache-Control'] = 'no-store'
    return response


def job_details(request):
    """Read one public advert on demand, or analyse user-pasted full text."""
    if request.method != 'POST':
        return JsonResponse({'error':'Method not allowed'},status=405)
    try:
        data=json.loads(request.body)
        role=data.get('role'); profile=data.get('profile')
        if not isinstance(role,dict) or not isinstance(profile,dict) or not profile.get('gradYear'):
            raise ValueError('Review your profile and select a role first.')
        matcher.stage_for(profile['gradYear'])
        pasted=data.get('text')
        if pasted is not None:
            if not isinstance(pasted,str) or not 150 <= len(pasted.strip()) <= 30000:
                raise ValueError('Paste between 150 and 30,000 characters of the full job description.')
            details={'text':pasted.strip(), 'source':'pasted job description', 'url':role.get('applyUrl','')}
        else:
            from .job_details import cached_advert, read_and_cache_advert
            details=cached_advert(role.get('applyUrl',''),role.get('title',''))
            if details and details.get('unavailable'):
                raise ValueError('Cached unavailable advert')
            if details is None:
                from .job_details import advert_read_capability
                capability=advert_read_capability()
                if not capability['fullAdvertReadsAvailable']:
                    return JsonResponse({'error':capability['fullAdvertReadNotice'],'code':'hosting_network_restricted'},status=503)
                limit=max(0,int(os.getenv('HIREATIONSHIP_DAILY_JOB_SEARCH_LIMIT','30')))
                if not reserve_live_search(request.user,limit):
                    return JsonResponse({'error':'Daily advert-reading/search allowance reached. You can still paste the description.'},status=429)
                details=read_and_cache_advert(role.get('applyUrl',''),role.get('title',''))
        advert={**role, 'summary':details['text'], 'fullDescription':True}
        assessments=claude.assess_live_jobs(profile,[advert]) if claude.has_key() else demo.assess_live_jobs(profile,[advert])
        assessment=next((a for a in assessments if a.get('id')==role.get('id')), {})
        evaluation=None
        outcome='unverified'
        reasons=assessment.get('reasons') or ['The description was read, but the current checker cannot confidently assess every requirement.']
        if assessment.get('eligible') is False:
            outcome='ineligible'
        elif assessment.get('eligible') is True and isinstance(assessment.get('role'),dict):
            extracted=assessment['role']
            evaluated_role={**role, **extracted, 'id':role['id'], 'summary':details['text'],
                'source':'Adzuna','eligibilityUnverified':False,'descriptionSource':details['source'],
                'matchingMode':'claude' if claude.has_key() else 'rules'}
            evaluation=matcher.evaluate(profile,evaluated_role)
            outcome='matched' if evaluation['tier']=='eligible' else 'ineligible'
            reasons=evaluation['hardFails']+(['Missing required skills: '+', '.join(evaluation['missingRequired'])] if evaluation['missingRequired'] else []) or assessment.get('reasons',[])
        mentioned=[skill for skill in demo.KNOWN_SKILLS if matcher.has_skill({'skills':[]}, skill, details['text'].lower())]
        supported=[skill for skill in mentioned if matcher.has_skill(profile,skill,matcher.haystack(profile))]
        result=JsonResponse({'mentionedSkills':mentioned,'supportedSkills':supported,'description':details['text'],'descriptionSource':details['source'],'sourceUrl':details['url'],
            'outcome':outcome,'reasons':reasons,'evaluation':evaluation if outcome=='matched' else None})
        result['Cache-Control']='no-store'
        return result
    except (ValueError,TypeError,KeyError,AttributeError):
        # Never echo request data, URLs, or upstream exceptions containing credentials.
        return JsonResponse({'error':'Could not read or assess this advert. Open the original advert and paste its full description (150–30,000 characters), or retry later.'},status=400)
